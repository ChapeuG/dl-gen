"""Leitor de schema — qualquer formato que traga o nome e o tipo dos campos vira DDL (CREATE TABLE).

O resto do framework continua lendo DDL (ddl_parser.py); este módulo só converte a entrada:

    DDL            CREATE TABLE ...                          (passa direto)
    Query          SELECT / WITH / CREATE VIEW ... AS SELECT  (tipo por CAST(x AS t), x::t ou CONVERT(t, x))
    JSON           Avro (.avsc), JSON Schema, StructType do Spark, lista de {name, type}
    Lista de campos  CSV/TSV/;/| ou planilha (.xlsx) com colunas nome + tipo; saída de DESCRIBE
    Arquivo de dados .parquet (schema do arquivo), .csv/.json com registros (tipos inferidos)

Quem não tem tipo (coluna de query sem CAST) vira varchar e ganha um comentário avisando.
"""

from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass, field
from pathlib import Path

NO_TYPE_COMMENT = "tipo não informado na origem (assumido texto)"


@dataclass
class Column:
    name: str
    type: str = "varchar"
    nullable: bool = True
    pk: bool = False
    comment: str = ""


@dataclass
class Table:
    name: str
    columns: list[Column] = field(default_factory=list)
    comment: str = ""


# ── Tipos de outros formatos → tipo SQL que o ddl_parser entende ──────

_TYPE_ALIASES = {
    "string": "varchar", "str": "varchar", "utf8": "varchar", "large_string": "varchar", "large_utf8": "varchar",
    "object": "varchar", "enum": "varchar", "fixed": "bytea", "char": "char",
    "long": "bigint", "int64": "bigint", "uint32": "bigint", "uint64": "decimal(20,0)", "integer": "int",
    "int32": "int", "uint16": "int", "short": "smallint", "int16": "smallint", "byte": "smallint", "int8": "smallint",
    "uint8": "smallint", "float": "real", "float32": "real", "halffloat": "real", "float16": "real",
    "float64": "double", "bool": "boolean", "date32": "date", "date64": "date", "date32[day]": "date",
    "binary": "bytea", "bytes": "bytea", "large_binary": "bytea", "null": "varchar",
}


def sql_type(raw: str) -> str:
    """Tipo de Avro, Spark, Arrow, pandas ou JSON → tipo SQL (varchar, bigint, decimal(p,s), timestamp...)."""
    t = str(raw or "").strip()
    low = t.lower()
    if not low:
        return "varchar"
    m = re.match(r"(?:array|list|large_list)\s*<\s*(?:item\s*:\s*)?(.+)>$", low)
    if m:
        return f"{sql_type(m.group(1))}[]"
    if re.match(r"(?:struct|map)\s*<", low) or low in ("struct", "map", "record"):
        return "json"
    m = re.match(r"decimal(?:128|256)?\s*\(\s*(\d+)\s*,\s*(\d+)\s*\)$", low)
    if m:
        return f"decimal({m.group(1)},{m.group(2)})"
    if low.startswith(("timestamp", "datetime")) and not low.startswith("datetime2"):
        return "timestamp"
    if low.startswith("time64") or low.startswith("time32"):
        return "time"
    return _TYPE_ALIASES.get(low, t)


def _clean_name(name: str) -> str:
    """Nome de coluna válido no DDL: sem aspas, espaços e símbolos viram _."""
    cleaned = re.sub(r"\W+", "_", str(name).strip().strip('"`[]'), flags=re.UNICODE).strip("_")
    return cleaned or "coluna"


def _sql_literal(text: str) -> str:
    return str(text or "").replace("'", '"').replace("\n", " ").strip()


def to_ddl(tables: list[Table]) -> str:
    """Tabelas → script DDL (um CREATE TABLE por tabela)."""
    if not tables or not any(t.columns for t in tables):
        raise ValueError("Nenhum campo encontrado (preciso do nome e do tipo de cada campo)")
    scripts = []
    for t in tables:
        if not t.columns:
            continue
        lines = []
        for c in t.columns:
            line = f"  {_clean_name(c.name)} {c.type or 'varchar'}"
            if not c.nullable:
                line += " NOT NULL"
            if c.comment:
                line += f" COMMENT '{_sql_literal(c.comment)}'"
            lines.append(line)
        pks = [_clean_name(c.name) for c in t.columns if c.pk]
        if pks:
            lines.append(f"  PRIMARY KEY ({', '.join(pks)})")
        tail = f" COMMENT '{_sql_literal(t.comment)}'" if t.comment else ""
        table_name = ".".join(_clean_name(part) for part in t.name.split("."))  # schema.tabela vira a origem
        scripts.append(f"CREATE TABLE {table_name} (\n" + ",\n".join(lines) + f"\n){tail};")
    return "\n\n".join(scripts) + "\n"


# ── Query (SELECT) ─────────────────────────────────────────────────────

_CTAS_RE = re.compile(r"create\s+(?:or\s+replace\s+)?(?:materialized\s+)?(?:view|table)\s+(?:if\s+not\s+exists\s+)?"
                      r"([^\s(]+)(?:\s*\([^)]*\))?\s+as\s*\(?\s*(?=select|with)", re.IGNORECASE)
_QUERY_START_RE = re.compile(r"^\s*\(?\s*(select|with)\b", re.IGNORECASE)


def _strip_sql_comments(sql: str) -> str:
    return re.sub(r"/\*.*?\*/", " ", re.sub(r"--[^\n]*", "", sql), flags=re.DOTALL)


def _top_level_keyword(sql: str, keyword: str, start: int = 0) -> int:
    """Posição da palavra-chave fora de parênteses e aspas (-1 se não houver)."""
    depth, quote = 0, ""
    pattern = re.compile(rf"\b{keyword}\b", re.IGNORECASE)
    i = start
    while i < len(sql):
        ch = sql[i]
        if quote:
            if ch == quote:
                quote = ""
        elif ch in ("'", '"', "`"):
            quote = ch
        elif ch == "(":
            depth += 1
        elif ch == ")":
            depth -= 1
        elif depth == 0 and pattern.match(sql, i) and (i == 0 or not (sql[i - 1].isalnum() or sql[i - 1] == "_")):
            return i
        i += 1
    return -1


def _split_commas(text: str) -> list[str]:
    from framework.parsers.ddl_parser import _split_top_level

    return [p.strip() for p in _split_top_level(text) if p.strip()]


def _main_select(sql: str) -> tuple[str, str]:
    """Lista de colunas e o resto (FROM ...) do SELECT principal — o último no nível de fora (depois dos CTEs)."""
    pos, last = 0, -1
    while True:
        found = _top_level_keyword(sql, "select", pos)
        if found < 0:
            break
        last, pos = found, found + 6
    if last < 0:
        raise ValueError("Query sem SELECT")
    body = sql[last + 6:]
    end = _top_level_keyword(body, "from")
    if end < 0:
        end = len(body)
    columns = re.sub(r"^\s*(?:distinct|all)\b|^\s*top\s+\d+", "", body[:end], flags=re.IGNORECASE)
    return columns, body[end:]


_CAST_RE = re.compile(r"^(?:try_)?cast\s*\((.+)\s+as\s+([\w ]+(?:\s*\([^)]*\))?(?:\[\])?)\s*\)$", re.IGNORECASE | re.DOTALL)
_CONVERT_RE = re.compile(r"^convert\s*\(\s*([\w ]+(?:\s*\([^)]*\))?)\s*,(.+)\)$", re.IGNORECASE | re.DOTALL)
_PG_CAST_RE = re.compile(r"^(.+?)::\s*([\w ]+(?:\s*\([^)]*\))?(?:\[\])?)$", re.DOTALL)
_AS_ALIAS_RE = re.compile(r"^(.*\S)\s+as\s+([\"`\[]?\w+[\"`\]]?)$", re.IGNORECASE | re.DOTALL)
_IMPLICIT_ALIAS_RE = re.compile(r"^(.*[\w)\]\"'`])\s+([A-Za-z_]\w*)$", re.DOTALL)
_NOT_ALIAS = {"end", "asc", "desc", "null", "distinct"}
_IDENT_RE = re.compile(r"[\w.\"`\[\]]+")
_FUNC_TYPES = {"count": "bigint", "current_date": "date", "current_timestamp": "timestamp", "now": "timestamp",
               "getdate": "timestamp", "sysdate": "timestamp", "to_date": "date", "to_timestamp": "timestamp"}


def _expr_type(expr: str) -> str:
    expr = expr.strip()
    for regex, type_group in ((_CAST_RE, 2), (_CONVERT_RE, 1), (_PG_CAST_RE, 2)):
        m = regex.match(expr)
        if m:
            return m.group(type_group).strip()
    if re.fullmatch(r"-?\d+", expr):
        return "int"
    if re.fullmatch(r"-?\d+\.\d+", expr):
        return "double"
    if re.fullmatch(r"'[^']*'", expr):
        return "varchar"
    if expr.lower() in ("true", "false"):
        return "boolean"
    m = re.match(r"(\w+)\s*(?:\(|$)", expr)
    return _FUNC_TYPES.get(m.group(1).lower(), "") if m else ""


def _column_from_select_item(item: str, position: int) -> Column:
    if item == "*" or item.endswith(".*"):
        raise ValueError("A query usa '*': liste as colunas (com CAST para informar o tipo)")
    expr, name = item, ""
    m = _AS_ALIAS_RE.match(item) or _IMPLICIT_ALIAS_RE.match(item)
    if m and m.group(2).lower() not in _NOT_ALIAS:
        expr, name = m.group(1).strip(), m.group(2)
    if not name:
        cast = _PG_CAST_RE.match(expr) or _CAST_RE.match(expr)
        base = (cast.group(1) if cast else expr).strip()
        name = base.split(".")[-1] if _IDENT_RE.fullmatch(base) else f"coluna_{position}"
    typ = _expr_type(expr)
    return Column(name=_clean_name(name), type=typ or "varchar", comment="" if typ else NO_TYPE_COMMENT)


def parse_query(sql: str, table: str = "") -> Table:
    """SELECT (ou CREATE VIEW/TABLE ... AS SELECT) → tabela com uma coluna por item do SELECT."""
    sql = _strip_sql_comments(sql).strip().rstrip(";")
    name = table
    m = _CTAS_RE.search(sql)
    if m:
        name = name or m.group(1)
        sql = sql[m.end():]
    columns_sql, rest = _main_select(sql)
    if not name:
        fm = re.match(r"\s*from\s+([\w.\"`\[\]]+)", rest, re.IGNORECASE)
        name = fm.group(1) if fm and not fm.group(1).startswith("(") else "consulta"
    columns = [_column_from_select_item(item, i) for i, item in enumerate(_split_commas(columns_sql), start=1)]
    return Table(name=name, columns=columns)


# ── JSON ───────────────────────────────────────────────────────────────

def _avro_type(t) -> tuple[str, bool]:
    """Tipo Avro → (tipo SQL, nullable)."""
    if isinstance(t, list):
        non_null = [x for x in t if x != "null"]
        typ, _ = _avro_type(non_null[0] if len(non_null) == 1 else "string")
        return typ, "null" in t
    if isinstance(t, dict):
        logical = t.get("logicalType", "")
        if logical == "decimal":
            return f"decimal({t.get('precision', 38)},{t.get('scale', 0)})", False
        if logical == "date":
            return "date", False
        if logical.startswith("timestamp") or logical.startswith("local-timestamp"):
            return "timestamp", False
        if logical == "uuid":
            return "uuid", False
        if t.get("type") == "array":
            return f"{_avro_type(t.get('items', 'string'))[0]}[]", False
        if t.get("type") in ("record", "map"):
            return "json", False
        return _avro_type(t.get("type", "string"))
    return sql_type(t), False


def _from_avro(doc: dict) -> Table:
    cols = []
    for f in doc.get("fields", []):
        typ, nullable = _avro_type(f.get("type", "string"))
        cols.append(Column(f["name"], typ, nullable, comment=f.get("doc", "")))
    return Table(doc.get("name", ""), cols, doc.get("doc", ""))


def _spark_type(t) -> str:
    if isinstance(t, dict):
        if t.get("type") == "array":
            return f"{_spark_type(t.get('elementType', 'string'))}[]"
        return "json"
    return sql_type(t)


def _from_struct(doc: dict, name: str = "") -> Table:
    cols = [Column(f["name"], _spark_type(f.get("type", "string")), bool(f.get("nullable", True)),
                   comment=(f.get("metadata") or {}).get("comment", "")) for f in doc.get("fields", [])]
    return Table(name, cols)


def _json_schema_type(prop: dict) -> tuple[str, bool]:
    t = prop.get("type", "string")
    nullable = False
    if isinstance(t, list):
        nullable = "null" in t
        t = next((x for x in t if x != "null"), "string")
    fmt = prop.get("format", "")
    if t == "string":
        if fmt == "date-time":
            return "timestamp", nullable
        if fmt == "date":
            return "date", nullable
        if fmt == "uuid":
            return "uuid", nullable
        return (f"varchar({prop['maxLength']})" if prop.get("maxLength") else "varchar"), nullable
    if t == "integer":
        return "bigint", nullable
    if t == "number":
        return "double", nullable
    if t == "array":
        return f"{_json_schema_type(prop.get('items') or {})[0]}[]", nullable
    if t == "object":
        return "json", nullable
    return sql_type(t), nullable


def _from_json_schema(doc: dict) -> Table:
    required = set(doc.get("required") or [])
    cols = []
    for name, prop in (doc.get("properties") or {}).items():
        typ, nullable = _json_schema_type(prop or {})
        cols.append(Column(name, typ, nullable or name not in required, comment=(prop or {}).get("description", "")))
    return Table(doc.get("title", ""), cols, doc.get("description", ""))


# Papel de cada coluna da lista de campos, pelo cabeçalho (normalizado: minúsculo, sem acento, com _).
# A ordem importa: "descricao_do_campo" é comentário, não nome.
_HEADER_ROLES = (
    ("type", ("tipo", "type")),
    ("comment", ("descri", "coment", "comment", "doc", "defini")),
    ("required", ("obrigat", "required", "not_null")),
    ("nullable", ("nul",)),
    ("pk", ("pk", "chave", "primary")),
    ("name", ("nome", "name", "campo", "coluna", "column", "col", "field", "atributo")),
)


def _header_role(header: str) -> str:
    key = _norm_key(header)
    return next((role for role, hints in _HEADER_ROLES if any(h in key for h in hints)), "")


def _norm_key(key) -> str:
    import unicodedata

    text = unicodedata.normalize("NFKD", str(key or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")


def _truthy(value) -> bool:
    return str(value).strip().lower() in ("true", "1", "s", "sim", "y", "yes", "x", "pk")


def _from_field_rows(rows: list[dict], name: str = "") -> Table | None:
    """Linhas com nome + tipo (lista JSON, CSV de campos, planilha). None se não tiver as duas colunas."""
    rows = [r for r in rows if isinstance(r, dict)]
    if not rows:
        return None
    roles: dict[str, str] = {}
    for header in rows[0]:
        role = _header_role(header)
        if role and role not in roles:
            roles[role] = header
    if "name" not in roles or "type" not in roles:
        return None

    def get(row: dict, role: str):
        value = row.get(roles[role]) if role in roles else None
        return None if value is None or str(value).strip() == "" else value

    cols = []
    for r in rows:
        col_name, typ = get(r, "name"), get(r, "type")
        if not col_name or str(col_name).startswith("#"):  # DESCRIBE: "# Partition Information"
            continue
        nullable = True
        if get(r, "nullable") is not None:
            nullable = _truthy(get(r, "nullable"))
        elif get(r, "required") is not None:
            nullable = not _truthy(get(r, "required"))
        pk = _truthy(get(r, "pk")) if get(r, "pk") is not None else False
        cols.append(Column(str(col_name).strip(), sql_type(str(typ or "")), nullable and not pk, pk,
                           str(get(r, "comment") or "").strip()))
    return Table(name, cols)


def _from_json(doc, name: str) -> list[Table]:
    if isinstance(doc, list):
        if doc and all(isinstance(d, dict) and d.get("type") == "record" for d in doc):
            return [_from_avro(d) for d in doc]
        table = _from_field_rows(doc, name)
        if table:
            return [table]
        if doc and all(isinstance(d, dict) for d in doc):
            return [_from_records(doc, name)]
        raise ValueError("JSON em lista sem campos reconhecíveis (esperado [{name, type}] ou registros de dados)")
    if not isinstance(doc, dict):
        raise ValueError("JSON sem schema reconhecível")
    if doc.get("type") == "record":
        return [_from_avro(doc)]
    if doc.get("type") == "struct":
        return [_from_struct(doc, name)]
    if "properties" in doc:
        return [_from_json_schema(doc)]
    for key in ("columns", "fields", "campos", "colunas", "schema"):
        if isinstance(doc.get(key), list):
            table = _from_field_rows(doc[key], doc.get("table") or doc.get("tabela") or doc.get("name") or name)
            if table:
                table.comment = doc.get("description") or doc.get("comment") or ""
                return [table]
    if isinstance(doc.get("tables"), list):
        return [t for d in doc["tables"] for t in _from_json(d, name)]
    return [_from_records([doc], name)]


# ── Dados (amostra) e lista de campos em texto ─────────────────────────

def _pandas_type(series) -> str:
    kind = str(series.dtype)
    if kind.startswith(("int", "Int")):
        return "bigint"
    if kind.startswith(("float", "Float")):
        return "double"
    if kind in ("bool", "boolean"):
        return "boolean"
    if kind.startswith("datetime"):
        return "timestamp"
    values = series.dropna().astype(str)
    if len(values) and values.str.fullmatch(r"\d{4}-\d{2}-\d{2}").all():
        return "date"
    if len(values) and values.str.fullmatch(r"\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}(:\d{2}(\.\d+)?)?(Z|[+-]\d{2}:?\d{2})?").all():
        return "timestamp"
    if series.dropna().map(lambda v: isinstance(v, (list, dict))).any():
        return "json"
    return "varchar"


def _from_frame(frame, name: str) -> Table:
    # Uma amostra não prova NOT NULL: todas as colunas ficam anuláveis
    return Table(name, [Column(str(c), _pandas_type(frame[c])) for c in frame.columns])


def _from_records(records: list[dict], name: str) -> Table:
    import pandas as pd

    return _from_frame(pd.DataFrame.from_records(records), name)


def _sniff_rows(text: str) -> list[dict] | None:
    """Texto com cabeçalho (CSV, TSV, ; ou |) → linhas como dict. Saída de DESCRIBE (espaços) também."""
    lines = [ln for ln in text.splitlines() if ln.strip() and not re.fullmatch(r"[\s|+=-]+", ln)]
    if not lines:
        return None
    header = lines[0]
    delimiter = next((d for d in ("\t", ";", "|", ",") if d in header), None)
    if delimiter:
        reader = csv.DictReader(io.StringIO("\n".join(lines)), delimiter=delimiter)
        rows = [{(k or "").strip(): (v or "").strip() for k, v in r.items() if isinstance(v, str)} for r in reader]
        if delimiter == "|":  # tabela em Markdown: | campo | tipo |
            rows = [{k: v for k, v in r.items() if k} for r in rows]
        if delimiter == "\t" and rows and not any(_header_role(k) == "name" for k in rows[0]):
            return _positional([ln.split("\t", 2) for ln in lines])  # DESCRIBE sem cabeçalho
        return rows
    # Sem delimitador: "nome tipo [comentário...]" (DESCRIBE, lista simples); 2+ espaços separam tipos com espaço
    split = [re.split(r"\s{2,}", ln.strip(), maxsplit=2) if re.search(r"\s{2,}", ln.strip()) else ln.split(None, 2)
             for ln in lines]
    return _positional(split, strict=True) if all(len(s) >= 2 for s in split) else None


def _known_type(raw: str) -> bool:
    from framework.parsers.ddl_parser import SPARK_TYPE_MAP
    from framework.standards.sources import SOURCES
    from framework.standards.tipagem import TIPAGEM_OFICIAL, normalize_sql_type

    base = normalize_sql_type(sql_type(raw)).removesuffix("[]")
    return base in {*SPARK_TYPE_MAP, *TIPAGEM_OFICIAL, *(k for s in SOURCES.values() for k in s.type_overrides)}


def _positional(split: list[list[str]], strict: bool = False) -> list[dict] | None:
    """Linhas sem cabeçalho reconhecido: 1ª coluna = nome, 2ª = tipo, 3ª = comentário.

    strict (texto solto, separado por espaços): só aceita se todo tipo for conhecido, para não ler qualquer frase.
    """
    has_header = _header_role(split[0][0]) == "name" and _header_role(split[0][1]) == "type"
    body = split[1:] if has_header else split
    if strict and not all(_known_type(s[1]) for s in body):
        return None
    return [{"name": s[0], "type": s[1], "comment": s[2] if len(s) > 2 else ""} for s in body]


def _read_xlsx(data: bytes, name: str) -> list[Table]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            rows = [r for r in ws.iter_rows(values_only=True) if r and any(c is not None for c in r)]
            if len(rows) < 2:
                continue
            header = [str(c or "").strip() for c in rows[0]]
            table = _from_field_rows([dict(zip(header, r)) for r in rows[1:]], name)
            if table:
                return [table]
    finally:
        wb.close()
    raise ValueError("Planilha sem colunas de nome e tipo do campo (ex: 'Campo' e 'Tipo')")


def _read_parquet(data: bytes, name: str) -> list[Table]:
    import pyarrow.parquet as pq

    schema = pq.read_schema(io.BytesIO(data))
    return [Table(name, [Column(f.name, sql_type(str(f.type)), f.nullable) for f in schema])]


# ── Entrada ────────────────────────────────────────────────────────────

FORMATS = ("ddl", "query", "json", "lista", "xlsx", "parquet", "csv")


def detect_format(text: str, filename: str = "") -> str:
    ext = Path(filename).suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        return "xlsx"
    if ext == ".parquet":
        return "parquet"
    stripped = _strip_sql_comments(text).strip()
    if stripped[:1] in ("{", "["):
        return "json"
    if _CTAS_RE.search(stripped) or _QUERY_START_RE.match(stripped):
        return "query"
    if re.search(r"create\s+table", stripped, re.IGNORECASE):
        return "ddl"
    return "lista"


def read_tables(source: str | bytes | Path, filename: str = "", table: str = "") -> tuple[str, list[Table] | None]:
    """Lê a entrada e devolve (formato, tabelas). Para DDL, tabelas = None (o texto já é o DDL)."""
    if isinstance(source, Path):
        filename = filename or source.name
        source = source.read_bytes()
    name = table or (Path(filename).stem.split(".")[0] if filename else "") or "tabela"
    fmt = detect_format("" if isinstance(source, bytes) and Path(filename).suffix.lower() in (".xlsx", ".xlsm", ".parquet")
                        else (source.decode("utf-8-sig") if isinstance(source, bytes) else source), filename)
    if fmt == "xlsx":
        return fmt, _read_xlsx(source if isinstance(source, bytes) else source.encode(), name)
    if fmt == "parquet":
        return fmt, _read_parquet(source if isinstance(source, bytes) else source.encode(), name)

    text = source.decode("utf-8-sig") if isinstance(source, bytes) else source
    if fmt == "ddl":
        return fmt, None
    if fmt == "query":
        return fmt, [parse_query(text, table)]
    if fmt == "json":
        try:
            doc = json.loads(text)
        except json.JSONDecodeError:  # JSON Lines (um registro por linha)
            doc = [json.loads(ln) for ln in text.splitlines() if ln.strip()]
        tables = _from_json(doc, name)
        for t in tables:
            t.name = table or t.name or name
        return fmt, tables
    rows = _sniff_rows(text)
    found = _from_field_rows(rows, name) if rows else None
    if found:
        return "lista", [found]
    if rows and Path(filename).suffix.lower() in (".csv", ".tsv", ".txt") or (rows and "," in text.splitlines()[0]):
        import pandas as pd

        delimiter = next((d for d in ("\t", ";", "|", ",") if d in text.splitlines()[0]), ",")
        return "csv", [_from_frame(pd.read_csv(io.StringIO(text), sep=delimiter), name)]
    raise ValueError("Formato não reconhecido. Use DDL (CREATE TABLE), uma query SELECT, JSON (Avro, JSON Schema, "
                     "StructType), planilha/CSV com as colunas nome e tipo, ou um arquivo de dados (.parquet, .csv)")


def read_schema(source: str | bytes | Path, filename: str = "", table: str = "") -> str:
    """Qualquer entrada com nome e tipo dos campos → script DDL (CREATE TABLE) para o ddl_parser."""
    fmt, tables = read_tables(source, filename, table)
    if tables is None:
        return source.decode("utf-8-sig") if isinstance(source, bytes) else (
            source.read_text(encoding="utf-8") if isinstance(source, Path) else source)
    return to_ddl(tables)
