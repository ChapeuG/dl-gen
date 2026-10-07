"""Tipagem oficial: tipo SQL de origem → tipo Spark (PySpark).

Camadas (a última prevalece):
  1. SPARK_TYPE_MAP do parser (tipos não cobertos pela tabela oficial: bigint, text, double...)
  2. TIPAGEM_OFICIAL — tabela oficial "Tipo Origem → Tipo Final"
  3. planilha Tipagem.xlsx informada via --tipagem (opcional)

Todo tipo final precisa estar em SPARK_TYPES (tipos suportados pelo Spark).
"""

from __future__ import annotations

import re
from pathlib import Path

# Tabela oficial "Tipo Origem → Tipo Final"
TIPAGEM_OFICIAL: dict[str, str] = {
    "uuid": "StringType",
    "varchar": "StringType",
    "timestamp": "TimestampType",
    "int": "IntegerType",
    "numeric": "DecimalType",
    "bool": "BooleanType",
    "date": "DateType",
    "datetime": "TimestampType",
}

# Tipos Spark permitidos como dataType
SPARK_TYPES: dict[str, str] = {
    # Numeric
    "ByteType": "Numeric", "ShortType": "Numeric", "IntegerType": "Numeric", "LongType": "Numeric",
    "FloatType": "Numeric", "DoubleType": "Numeric", "DecimalType": "Numeric",
    # String / Binary / Boolean
    "StringType": "String", "BinaryType": "Binary", "BooleanType": "Boolean",
    # Datetime
    "DateType": "Datetime", "TimestampType": "Datetime",
    # Complex
    "ArrayType": "Complex", "MapType": "Complex", "StructType": "Complex", "StructField": "Complex",
}

_SOURCE_HEADERS = ("origem", "sql", "source", "banco", "ddl")
_TARGET_HEADERS = ("final", "scala", "spark", "destino", "target", "data lake", "datalake")


def base_spark_type(spark_type: str) -> str:
    """DecimalType(10,2) → DecimalType ; ArrayType(StringType) → ArrayType"""
    return spark_type.split("(", 1)[0].strip()


def is_valid_spark_type(spark_type: str) -> bool:
    if base_spark_type(spark_type) not in SPARK_TYPES or base_spark_type(spark_type) == "StructField":
        return False
    m = re.match(r"\w+\((.*)\)$", spark_type)
    if m and base_spark_type(spark_type) == "ArrayType":
        return is_valid_spark_type(m.group(1).split(",")[0].strip())
    return True


def _find_columns(header: tuple) -> tuple[int, int] | None:
    cells = [str(c or "").strip().lower() for c in header]
    src = next((i for i, c in enumerate(cells) if any(h in c for h in _SOURCE_HEADERS)), None)
    dst = next((i for i, c in enumerate(cells) if i != src and any(h in c for h in _TARGET_HEADERS)), None)
    if src is None or dst is None:
        return None
    return src, dst


def normalize_sql_type(raw_type: str) -> str:
    """VARCHAR(255) → varchar ; timestamp without time zone → timestamp without time zone"""
    return re.sub(r"\s+", " ", re.sub(r"\(.*?\)", "", raw_type.lower())).strip()


def load_type_map(path: str | Path) -> dict[str, str]:
    """Lê a planilha e retorna {tipo_sql_normalizado: TipoSpark}.

    Cabeçalho esperado: "Tipo Origem" / "Tipo Final" (também aceita "Tipo SQL" / "Tipo Scala").
    Sem cabeçalho reconhecível, usa as duas primeiras colunas.
    """
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    chosen: tuple | None = None
    for ws in wb.worksheets:
        rows = [r for r in ws.iter_rows(values_only=True) if r and any(c is not None for c in r)]
        if not rows:
            continue
        cols = _find_columns(rows[0])
        if cols:
            chosen = (rows[1:], cols)
            break
        if chosen is None:
            chosen = (rows, (0, 1))
    wb.close()

    if chosen is None:
        raise ValueError(f"Planilha de tipagem vazia: {path}")

    rows, (src, dst) = chosen
    type_map: dict[str, str] = {}
    invalid: list[str] = []
    for row in rows:
        if len(row) <= max(src, dst) or row[src] is None or row[dst] is None:
            continue
        sql_type = normalize_sql_type(str(row[src]))
        spark_type = str(row[dst]).strip().removesuffix("()")
        if not sql_type or not spark_type:
            continue
        if not is_valid_spark_type(spark_type):
            invalid.append(f"{row[src]} → {spark_type}")
            continue
        type_map[sql_type] = spark_type

    if invalid:
        raise ValueError(f"Tipos inválidos em {path} (permitidos: {', '.join(SPARK_TYPES)}): {'; '.join(invalid)}")
    if not type_map:
        raise ValueError(f"Nenhum mapeamento Tipo Origem → Tipo Final encontrado em {path}")
    return type_map
